"""報告スライド（PPTX）の各ページの文字を書き出して docs/slides_text.md を作る。

文章の控えをスライドの実物と一致させるため、手で書かずに生成物から書き出す。
使い方: uv run --with python-pptx python scripts/dump_slide_text.py
"""
import re

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]

src = ROOT / "reports" / "slides" / "ett_oil_temperature_poc.pptx"
out = ROOT / "docs" / "slides_text.md"
prs = Presentation(src)
lines = [
    "# 報告スライドの文章（v4・2026-09-23）",
    "",
    "`reports/slides/ett_oil_temperature_poc.pptx` から各ページの文字を書き出したもの（`scripts/build_slides.js` が生成元。数値は `reports/slide_data.json`）。",
    "v4 は v3（学術・コンサル型の作法）をもとに文字を減らした版。タイトルを 1 行にし、グラフのページの説明を 1〜2 項目に絞り、背景・評価のやり方・設計の判断・限界・結論・付録 D を図か表に置き換え、予測と実測を重ねた 1 週間の図を足した。要旨・問い・安定性の 3 枚は v3 のまま。",
    "グラフの目盛りや系列名は省いた。グラフの中の数値（棒の値・折れ線）は `reports/slide_data.json` を参照。",
    "",
]
for i, slide in enumerate(prs.slides, 1):
    texts, tables = [], []
    for sh in sorted(slide.shapes, key=lambda s: (Emu(s.top or 0), Emu(s.left or 0))):
        if sh.has_text_frame:
            t = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
            t = t.strip()
            # グラフの目盛り・季節名・月などの短いラベルは省く
            if t and len(t) > 3 and not re.fullmatch(r"\d+月", t):
                texts.append(t)
        elif sh.has_table:
            rows = [[c.text.replace("\n", " ") for c in r.cells] for r in sh.table.rows]
            tables.append(rows)
    lines.append(f"## {i}")
    lines.append("")
    for t in texts:
        lines.append(t.replace("\n", "  \n"))
        lines.append("")
    for rows in tables:
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("|" + "---|" * len(rows[0]))
        for r in rows[1:]:
            lines.append("| " + " | ".join(r) + " |")
        lines.append("")
open(out, "w", encoding="utf-8").write("\n".join(lines).rstrip() + "\n")
print("wrote", out.relative_to(ROOT), len(prs.slides), "slides")
